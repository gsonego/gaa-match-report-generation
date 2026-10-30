"""Build outputs/pilot_error_table.xlsx (claim-level check of the 9 valid pilot reports).

The rows were written by hand after reading each report against pilot_inputs.json;
this script only turns them into the workbook, so the table can be regenerated or extended.
Run from the repo root:  python scripts/build_pilot_error_table.py
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
OK, ERR, PART = 'ok', 'error', 'partly'
NDT, NDU = 'not in data (true)', 'not in data (unverified)'
R = []
def add(g, claim, typ, verdict, etype='', note=''):
    R.append((g, claim, typ, verdict, etype, note))
WD = 'weekday is not in the input; computed and correct'
g = 1
add(g,'Mayo 2-14 Galway 2-8','final score',OK)
add(g,'Mayo won / claimed the Connacht title','winner',OK)
add(g,'six-point victory','margin',OK)
add(g,'Sunday, July 25','date',OK,'',WD)
add(g,"O'Donoghue and Ruane contributed 2-5 between them",'scorer total',OK)
add(g,"O'Donoghue 1-3, including a penalty and two pointed frees",'scorer detail',OK)
add(g,'Ruane 1-2','scorer total',OK)
add(g,'10 Mayo players on the scoresheet (repeated as "Mayo\'s 10")','count',ERR,'wrong count','the data lists 9 Mayo scorers')
add(g,'Loftus and Conroy two points each','scorer total',OK)
add(g,"Hennelly converted a '45",'scorer detail',OK)
add(g,'Hennelly is the goalkeeper','added info',NDT,'','position not in the input')
add(g,'Durcan, McHale, McLoughlin, Carr one point each','scorer total',OK)
add(g,'both teams scored two goals','final score',OK)
add(g,'14 to 8 in points accounted for the whole margin','derived',OK)
add(g,"Galway's leading scorer was Tierney, four points",'ranking',PART,'ranking','Walsh 1-1 is also 4 points (tie)')
add(g,'Tierney: three of his four from frees','scorer detail',OK)
add(g,'Walsh 1-1 and Comer a goal','scorer total',OK)
add(g,'Paul Conroy two, Sweeney one','scorer total',OK)
add(g,'Galway had five scorers','count',OK)
g = 2
add(g,'Laois 2-17 Wexford 2-13','final score',OK)
add(g,'Laois won, four-point victory','margin',OK)
add(g,'Sunday, April 9','date',OK,'',WD)
add(g,'Barry 2-4, two goals','scorer total',OK)
add(g,"O'Carroll six points",'scorer total',OK)
add(g,'Laois total of 23 points','derived',OK)
add(g,'Wexford matched Laois on goals','final score',OK)
add(g,'Brosnan 1-7','scorer total',OK)
add(g,'Barry and Brosnan ten points each','derived',OK)
add(g,"Barry's four points: two frees and a mark",'scorer detail',OK)
add(g,"O'Carroll: a free, a mark and a 45 among six",'scorer detail',OK)
add(g,'Kingston four, three from frees','scorer detail',OK)
add(g,"Larkin two, O'Sullivan one",'scorer total',OK)
add(g,'Brosnan: penalty and five frees among seven points','scorer detail',OK)
add(g,'Brooks the other goal plus a point','scorer total',OK)
add(g,"S Nolan two; Hughes and O'Grady on target; E Nolan a 45",'scorer total',OK)
add(g,"Barry, O'Carroll and Kingston all but three of 23 points",'derived',OK)
add(g,'"opening championship assignment" (preliminary round)','stage',NDT,'','not stated in the input; plausible')
g = 3
add(g,'Armagh 0-20 Antrim 1-8','final score',OK)
add(g,'Armagh won by nine points','margin',OK)
add(g,'Saturday, April 8','date',OK,'',WD)
add(g,'Turbitt eight points','scorer total',OK)
add(g,'Armagh had eight scorers','count',OK)
add(g,'Turbitt: three frees, a mark, four from play','scorer detail',OK,'','"from play" inferred from the absence of a note')
add(g,'McPartlan four, all from play','scorer detail',OK,'','"from play" inferred from the absence of a note')
add(g,'Grugan three, including a free','scorer detail',OK)
add(g,"the trio accounted for 15 of Armagh's 20 points",'derived',OK)
add(g,'Stewart the only goal, plus a point','scorer total',OK)
add(g,"Antrim's other seven points shared among seven players",'derived',OK)
add(g,'McCann (Creggan) and Murray one free each','scorer detail',OK)
add(g,"McCabe, McAuley, Shivers, McCann (Aghagallon), Eastwood one each",'scorer total',OK)
add(g,'Rafferty one free; Mackin, Campbell, Hall, Kelly one from play each','scorer detail',OK)
g = 4
add(g,'Derry 0-14 Monaghan 0-14, a draw','final score',OK)
add(g,'draw (no winner)','winner',OK)
add(g,'Saturday, May 27; Round 1','date',OK,'',WD)
add(g,'neither side scored a goal','final score',OK)
add(g,'McGuigan nine points','scorer total',OK)
add(g,'McGuigan: six frees, a mark, two from play','scorer detail',OK)
add(g,"McGuigan's placed balls were half of Derry's 14",'derived',OK,'','7 of 14')
add(g,'Toner two, one from a free; Doherty, McFaul, Murray one each','scorer total',OK)
add(g,'Monaghan had eight scorers','count',OK)
add(g,'McCarthy and Bannigan three each, one free each, six together','scorer total',OK)
add(g,'Beggan two frees','scorer detail',OK)
add(g,'Beggan is the goalkeeper','added info',NDT,'','position not in the input')
add(g,"Carey two; Wylie, O'Connell, Boyle, Mohan one each",'scorer total',OK)
g = 5
add(g,'Cork 1-14 Roscommon 0-16','final score',OK)
add(g,'Cork won by one point','margin',OK)
add(g,'Saturday, June 24','date',OK,'',WD)
add(g,'Cork secured a place in the quarter-finals','stage',NDT,'','progression not in the input; correct for a preliminary quarter-final')
add(g,'Murtagh brothers 11 points combined (6 + 5)','derived',OK)
add(g,'Diarmuid and Ciarain Murtagh are brothers','added info',NDT,'','not in the input')
add(g,'Sherlock five, four from frees','scorer detail',OK)
add(g,'Maguire and Jones two each','scorer total',OK)
add(g,"Corbett's goal the only goal",'scorer total',OK)
add(g,"D Murtagh six including three frees, the game's top scorer",'ranking',OK)
add(g,'C Murtagh five, three from frees','scorer detail',OK)
add(g,"Walsh, Fahy, Taylor, O'Driscoll, O'Donovan one each",'scorer total',OK)
add(g,'Cork had nine scorers','count',OK)
add(g,'Cregg and Smith two each; Cregg included a mark','scorer detail',OK)
add(g,"Carroll converted a '45",'scorer detail',OK)
add(g,'Carroll is the goalkeeper','added info',NDU,'','position not in the input; not checked')
add(g,"Cork's 1-14 amounted to 17 points against 16",'derived',OK)
add(g,'Roscommon had more individual scores (16 v 15)','derived',OK)
g = 6
add(g,'Louth 0-17 Kildare 0-13','final score',OK)
add(g,'Louth won by four points','margin',OK)
add(g,'Sunday, April 28','date',OK,'',WD)
add(g,'Louth into the Senior Football Championship final / "booked their passage to the decider" (headline and text)','stage',ERR,'wrong context','the input says only "Semi-final"; the model read it as the All-Ireland semi-final and invented a final')
add(g,'semi-final decided entirely by points','final score',OK)
add(g,'Mulroy seven points','scorer total',OK)
add(g,"Mulroy: two from play, four frees, a '45",'scorer detail',OK)
add(g,'Louth eight scorers','count',OK)
add(g,'Durnin, Grimes, Lennon two each','scorer total',OK)
add(g,'Burns, Downey, Mathews, Sharkey one each','scorer total',OK)
add(g,"Louth 12 points from play to Kildare's six",'derived',OK,'','assumes unmarked points are from play')
add(g,'Kildare 7 of 13 from placed balls; Louth 5 of 17','derived',OK)
add(g,'Feely top scorer, four, two from frees','scorer detail',OK)
add(g,'Farrell three frees; Woodgate two placed balls','scorer detail',OK)
add(g,'Kirwan two from play; Archbold and Flynn one each','scorer total',OK)
add(g,"Louth's other scorers 10 points",'derived',OK)
add(g,'first name "Sam" for Mulroy','added info',NDT,'','not in the input')
g = 7
add(g,'Dublin 0-17 Mayo 0-17, a draw','final score',OK)
add(g,'draw (no winner)','winner',OK)
add(g,'Sunday, June 16; Round 3','date',OK,'',WD)
add(g,'neither side scored a goal','final score',OK)
add(g,"Costello and O'Donoghue seven points each",'scorer total',OK)
add(g,'Costello four from play and three frees','scorer detail',OK)
add(g,"five of O'Donoghue's seven from placed balls",'scorer detail',OK)
add(g,'Ten Dublin scorers','count',OK)
add(g,'Basquel two points','scorer total',OK)
add(g,'"Brian Howard\'s fellow Dublin contributors" (Howard is not in the input)','invented entity',ERR,'invented name','no Howard among the Dublin scorers')
add(g,'Fenton was not among the scorers','scorer total',OK)
add(g,"Bugler, Kilkenny, Mannion, McCaffrey, Murchan, O'Callaghan, Scully, Small one each",'scorer total',OK)
add(g,'Mayo eight scorers','count',OK)
add(g,'Conroy three; Ruane two; Coen, Flynn, Loftus one each from play','scorer total',OK)
add(g,"O'Shea a free; Reape a 45",'scorer detail',OK)
add(g,'Reape is the goalkeeper','added info',NDT,'','position not in the input')
add(g,"Dublin 14 points from play to Mayo's 10",'derived',OK)
add(g,'first names Cormac, Ryan, Colm','added info',NDT,'','not in the input')
g = 8
add(g,'Donegal 2-23 Clare 0-5','final score',OK)
add(g,'Donegal won by 24 points','margin',OK)
add(g,'Saturday, June 15; Round 3','date',OK,'',WD)
add(g,'Gallen and McHugh both 1-4, the two goals','scorer total',OK)
add(g,'Donegal spread across 11 scorers','count',OK)
add(g,'Clare five points, four by McMahon','scorer total',OK)
add(g,"Donegal's final tally of 29 points",'derived',OK)
add(g,'McBrearty five including one free','scorer detail',OK)
add(g,'Gallen and McHugh seven points each','derived',OK)
add(g,'Gallen included two converted 45s','scorer detail',OK)
add(g,"C O'Donnell and Thompson two; six others one each",'scorer total',OK)
add(g,"three of McMahon's four from frees; O'Donnell the only other Clare score",'scorer detail',OK)
add(g,'Clare two points from play','derived',OK)
add(g,'two goals and 23 points','final score',OK)
g = 9
add(g,'Galway 0-17 Dublin 0-16','final score',OK)
add(g,'Galway won by one point','margin',OK)
add(g,'Saturday, June 29; quarter-final; semi-final place','date',OK,'',WD)
add(g,"Galway ended Dublin's reign as All-Ireland champions",'stage',NDT,'','not in the input; true (Dublin won the 2023 final 1-15 to 1-13, per the saved 2023 page)')
add(g,'Walsh seven: four frees, three from play','scorer detail',OK)
add(g,'McDaid three','scorer total',OK)
add(g,"Costello and O'Callaghan four each",'scorer total',OK)
add(g,"Costello: two frees and a '45; O'Callaghan: two marks and a free",'scorer detail',OK)
add(g,'Fenton and Mannion also scored','scorer total',OK)
add(g,'Bugler and Kilkenny two each; McGarry and Small one; Small from a mark','scorer total',OK)
add(g,'Tierney one from a mark','scorer detail',OK)
add(g,"\"Cillian O Curraoin's attacking colleagues Culhane and Heaney\" (O Curraoin is not in the input)",'invented entity',ERR,'invented name','no such player among the Galway scorers')
add(g,'Conroy, Culhane, Darcy, Heaney, Maher, McHugh one point each','scorer total',OK)
add(g,'first names for Dublin players (Brian Fenton, Paul Mannion, Sean Bugler, Ciaran Kilkenny)','added info',NDT,'','not in the input; Fenton and Mannion appear in full on the saved 2023 page')
add(g,'first names for Galway players (Cillian McDaid, Patrick Culhane, Johnny Heaney, Cein Darcy, John Maher, Dylan McHugh)','added info',NDU,'','not in the input; not checked')
dates = {1:'2021-07-25',2:'2023-04-09',3:'2023-04-08',4:'2023-05-27',5:'2023-06-24',6:'2024-04-28',7:'2024-06-16',8:'2024-06-15',9:'2024-06-29'}
wd = {1:'Sunday',2:'Sunday',3:'Saturday',4:'Saturday',5:'Saturday',6:'Sunday',7:'Sunday',8:'Saturday',9:'Saturday'}
for k, v in dates.items():
    assert dt.date.fromisoformat(v).strftime('%A') == wd[k], k

wb = Workbook(); ws = wb.active; ws.title = 'Claims'
ws.append(['#','Game','Claim (as written, shortened)','Claim type','Verdict','Error type','Note'])
for i, (g, c, t, v, e, n) in enumerate(R, 1):
    ws.append([i, g, c, t, v, e, n])
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496')
for col, w in zip('ABCDEFG', (5, 6, 70, 16, 24, 16, 60)):
    ws.column_dimensions[col].width = w
fills = {ERR: 'F8CBAD', PART: 'FFE699', NDU: 'DDEBF7', NDT: 'EDEDED'}
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    f = fills.get(row[4].value)
    if f:
        for c in row: c.fill = PatternFill('solid', fgColor=f)
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
dv = DataValidation(type='list', formula1='"ok,error,partly,not in data (true),not in data (unverified)"', allow_blank=False)
ws.add_data_validation(dv); dv.add(f'E2:E{len(R)+1}')

n = len(R) + 1
sm = wb.create_sheet('Summary')
sm.append(['Claim type','Claims','ok','error','partly','not in data (true)','not in data (unverified)'])
types = sorted({r[2] for r in R})
for i, t in enumerate(types, 2):
    sm.append([t, f'=COUNTIF(Claims!$D$2:$D${n},A{i})'] +
              [f'=COUNTIFS(Claims!$D$2:$D${n},$A{i},Claims!$E$2:$E${n},"{v}")' for v in ('ok','error','partly','not in data (true)','not in data (unverified)')])
k = len(types) + 2
sm.append(['Total'] + [f'=SUM({c}2:{c}{k-1})' for c in 'BCDEFG'])
for c in sm[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496')
for c in sm[k]: c.font = Font(bold=True)
sm.column_dimensions['A'].width = 18
for col in 'BCDEFG': sm.column_dimensions[col].width = 22

nt = wb.create_sheet('Notes')
for line in [
 'Pilot run 2026-10-28_2207, prompt v1 (condition A), model gpt-6.1-sol. Games 1-9 only: game 10 (Armagh v Kerry) went to extra time and is out of scope.',
 'Each claim was checked by hand against data/processed/pilot_inputs.json (scorers, score and the truth block).',
 'Verdicts: ok = supported by the input or follows from it; error = contradicts the input or invents something; partly = literally true but misleading; not in data = not in the input, so it comes from the model (true / unverified).',
 '"From play" claims are counted as ok when the input has no note on the points (the notation only marks frees, marks, 45s and penalties).',
 'The weekday of each date is not in the input; all nine were computed and match.',
 'The list is a hand selection of the factual claims of each report, not an automatic extraction. Interpretive phrases ("convincing", "deserved") are not claims.',
 'One reader only. A second pass over the "ok" rows would reduce the chance of a missed error.']:
    nt.append([line])
nt.column_dimensions['A'].width = 140
OUT = 'outputs/pilot_error_table.xlsx'
wb.save(OUT)
from collections import Counter
print(len(R), Counter(r[3] for r in R))
print(Counter((r[2], r[3]) for r in R if r[3] != 'ok'))
print(Counter(r[2] for r in R))
